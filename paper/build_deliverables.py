"""Reproducible personal-project PDF deck, editable PPTX and seven-page technical report."""
from pathlib import Path
import json,html,re
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.colors import HexColor,white
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import Paragraph,SimpleDocTemplate,PageBreak,Spacer,Table,TableStyle
from reportlab.lib.pagesizes import A4
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
ROOT=Path(__file__).resolve().parents[1];DEST=ROOT/'paper';FINAL=ROOT/'verification/final'
AUTHOR='Praveen Sudireddy';ROLE='Digital Specialist Engineer (DSE), Infosys';IDENTITY=AUTHOR+' | '+ROLE
for name,file in [('AG','DejaVuSans.ttf'),('AGB','DejaVuSans-Bold.ttf')]:pdfmetrics.registerFont(TTFont(name,'/usr/share/fonts/truetype/dejavu/'+file))
pdfmetrics.registerFontFamily('AG',normal='AG',bold='AGB',italic='AG',boldItalic='AGB')
NAVY='#102b46';TEAL='#008a83';SLATE='#475569';LIGHT='#eff5f6'
# Same text is used for the PDF slides and editable PowerPoint; no old company/founding-location claims.
SLIDES=[
 ('Independent personal project','AetherGrid','Cold-chain intervention software with explicit simulation boundaries',
  'Author and scope',[AUTHOR,ROLE,'Professional role is author context only; no employer sponsorship or endorsement is claimed.'],
  'What the project demonstrates',['Detect heuristic distress and qualify a compatible buyer.','Reserve simulated resources; require operator confirmation.','Record an auditable simulated ledger, not a physical rescue.']),
 ('Problem','An alert is not a completed intervention.','A qualitative engineering problem, not a market-size or national-loss claim.',
  'Questions to resolve',['Who can authorize a cargo redirect?','Is a buyer compatible, reachable and available for the full lot?','How are reservations, confirmation and receipt reconciled?'],
  'Project approach',['Connect telemetry, model state and reviewed transaction workflow.','Reject unsuitable bids and release expired reservations once.','Keep failed lots and uncertainty visible; never treat an award as delivery.']),
 ('Scope','A narrow, supervised mango use case.','Bengaluru–Hyderabad is a proposed pilot corridor, not an office or deployment claim.',
  'Proposed pilot boundary',['Mangoes; 5–10 supervised trucks only after access and consent.','Manual operator decisions with a no-money ledger.','Existing telemetry/log exports may be used; hardware fabrication is excluded.'],
  'Not established',['No observed field pilot, validated quality forecast or measured benefit.','No real bank escrow, UPI or insurer execution.','Legacy NH-44 map, buyers and carriers are demonstration data.']),
 ('Architecture','Five layers, with a visible provider boundary.','React portals → native ASP.NET workflow → SQLite documents and durable outbox.',
  'L0–L2',['L0: software telemetry provider and local simulators.','L1: software HMAC v2, identity binding, time and sequence checks.','L2: BioTimer, thermal Dijkstra, buyer eligibility and settlement arithmetic.'],
  'L3–L4',['L3: transactional reservations, reviewed settlement and reported receipt.','L4: Command, Buyer, Analytics, Carrier and Insurance portals.','Python remains a reference/demo implementation; native .NET uses HTTP + SSE.']),
 ('Lifecycle','A buyer reserves. An operator confirms.','Browser-tested with separate authenticated buyer and operator accounts.',
  'Initiation',['Record an operator-reviewed initial biological baseline.','Accept signed telemetry and open a time-boxed distress auction.','Check full lot, policy, capacity, cash, radius, reachability and deadline.'],
  'Completion controls',['Reserve simulated funds/capacity atomically, without settlement.','Operator confirms/rejects; timeouts release once with bounded fallback.','Reported full-lot receipt releases capacity once; disputes require review.']),
 ('Model','Q10 is a heuristic, not a food-safety forecast.','No trained ML model or independently validated prediction accuracy is claimed.',
  'BioTimer calculation',['Rate = Q10^((T−5)/10) × humidity factor × ethylene factor.','Consumed life += rate × exposure hours / reference life; door-open factor is 5.','Mango parameters: 264 h reference life, Q10 2.8, optimal 12°C.'],
  'Interpretation and limits',['Optimistic RVH differs from current-temperature spoilage ETA.','Native receiver persists checkpoints and uses signed elapsed time.','Parameters, maturity, handling and quality endpoints need real-data review.']),
 ('Worked example','The stated baseline gives 90.65% recovery.','Synthetic 1,000 kg mango lot; ₹120/kg contract baseline; ₹78/kg rescue bid.',
  'Explicit arithmetic',['Contract value ₹1,20,000; simulated salvage ₹78,000.','Sum assured ₹72,000; cover cap ₹32,400; deductible 5%.','Calculated insurance credit ₹30,780; buyer discount 35%.'],
  'Outcome classification',['Farmer simulated ledger credit: ₹1,08,780.','₹1,08,780 / ₹1,20,000 = 90.65%; not a universal guarantee.','No real payment, insurer acceptance or independent physical receipt.']),
 ('Routing','Rank delivery, liquidation and cold holding.','Reference-algorithm parity is distinct from validated navigation.',
  'Implemented native planning',['Bidirectional thermal Dijkstra over a fixed demonstration graph.','Costs include temperature, reefer state, coverage, road/risk and life urgency.','Operational buyer filters add policy, commodity, capacity and model reachability.'],
  'Decision boundaries',['70 cases / seven profiles: 2,100 fields matched the Python ranked plans.','Optional automatic model target selection never dispatches a physical truck.','Cold holding may slow decay; it never resets consumed biological life.']),
 ('Current verification','Executed checks, with their actual scope.','Fresh builds and isolated synthetic databases; results are not field efficacy.',
  'Regression and native services',['51 Python regression tests passed.','2,230 existing BioTimer/settlement/eligibility parity assertions passed.','Published native build, signed telemetry and transactional HTTP checks passed.'],
  'Real browser journey',['Buyer login → reserve; separate operator login → confirm.','UI shows ₹1,08,780; another buyer cannot read that award or totals.','Five portals load; receipt/duplicate release and CSV reporting checked through API.']),
 ('Accounting','Every rupee needs a stated baseline.','Deterministic simulation arithmetic is not insurance adjudication.',
  'Rules',['Insurance = min(shortfall, cover × assured amount) × (1−deductible).','Farmer total includes salvage, calculated insurance and bounded penalty top-up.','A target floor is not guaranteed when the available top-up is insufficient.'],
  'Evidence and reconciliation',['Confirmed settlements retain an evidence snapshot and hash.','Provider sandbox obligations use integer paise and explicit idempotency keys.','Terminal callback changes, wrong amounts and stale events are rejected.']),
 ('Security','Software integrity is not sensor truth.','No hardware signing, device mTLS or physical-tamper proof is implied.',
  'Implemented controls',['Opaque HttpOnly sessions; PBKDF2 hashes; role and buyer/device bindings.','CSRF, same-origin mutation checks, login throttling and request limits.','Persisted replay cursor, password recovery, session revocation and key rotation.'],
  'Boundary and remediation',['Acceptance review found and fixed a buyer-role portal access defect.','Buyer summary, auctions and settlements are now scoped to that identity.','Host compromise, external pentest, MFA/SSO and TLS deployment remain separate risks.']),
 ('Operations','Crash recovery must preserve obligations.','Forced-process termination was tested with a pending synthetic reservation.',
  'Durable behavior',['SQLite WAL transactions persist state with a sequenced event outbox.','SIGKILL/restart: pending resources released once after deadline.','Stale telemetry alerts and expired-authentication cleanup were exercised.'],
  'Safe operations',['Retain telemetry, ledger, evidence, callbacks and replay cursors.','Use online SQLite backup; rehearse restore before reopening ingestion.','A restored old backup may lose recent replay watermarks; reconcile before use.']),
 ('Project uses','Learning, rehearsal and auditable decision support.','Personal engineering work; no paid traction, commercial forecast or partnership claim.',
  'Useful applications',['Explore how telemetry can initiate a reviewed cold-chain workflow.','Study full-lot eligibility, reservation races and payout conservation.','Demonstrate React/.NET integration, identity controls and failure recovery.'],
  'Engineering learning',['Compare Python and C# reference calculations and routing outputs.','Rehearse no-money pilot procedures without building hardware.','Make assumptions, failure denominators and unverified outcomes explicit.']),
 ('No-access lab','Useful testing without provider credentials.','An invented local HTTP contract is not a named provider sandbox.',
  'Provider rehearsal',['Actual local HTTP timeout after commit and ambiguous 502 response tested.','Reconcile by lookup instead of blindly repeating financial submission.','Idempotency, callback duplication, wrong amount and reversal tests passed.'],
  'Calibration rehearsal',['Seven synthetic commodity pipelines require explicit synthetic opt-in.','Repeated lot IDs across train/holdout are rejected.','Candidate reports do not alter runtime Q10 or establish biological validity.']),
 ('Progress journal','From initial demo to reviewed software acceptance.','Work-package groupings, not invented calendar-week timesheets.',
  'Foundation and corrections',['Initial FastAPI/React demo and five portals; cold-start transport fixes.','Claims audit replaced unsupported recovery/field assertions.','Manual confirmation, eligibility and bounded fallback became explicit.'],
  'Hardening and latest work',['v2: access, checkpoints, calibration/pilot tools and calculation parity.','v3–v4: native workflow, auth, routing, recovery, offline lab and deployment checks.','Current: fresh two-role browser journey, identity correction and safe consolidation.']),
 ('Remaining gates','A tested prototype is not a certified deployment.','Hardware fabrication remains excluded, not a pending task.',
  'Prepared / not executed here',['Docker/Compose and container-build/image-scan CI are supplied.','No Docker runtime/daemon is available in this sandbox; image scan unexecuted.','Cloud/TLS deployment, external security review and operations sign-off remain open.'],
  'Requires external evidence',['Select a permitted provider product and obtain test access for live API validation.','Collect genuine controlled quality observations and independent holdout lots.','Obtain operator/owner consent, quality review and actual outcome evidence before a trial.']),
 ('Conclusion','An inspectable project—not a claimed physical rescue.',IDENTITY,
  'Delivered project artifacts',['Editable 17-slide PPTX and matching slide PDF.','Seven-page project technical/research report; no peer-review claim.','One master README, source, essential tests and retained evidence.'],
  'Read the results correctly',['Implemented: tested software paths and recorded state transitions.','Simulated: telemetry inputs, balances, insurance and farmer credits.','Unverified: real delivery benefit, biological accuracy and live-provider acceptance.'])]

def draw_paragraph(c,text,x,y,width,size=13,color=NAVY,bold=False):
 st=ParagraphStyle('p',fontName='AGB' if bold else 'AG',fontSize=size,leading=size*1.36,textColor=HexColor(color));p=Paragraph(html.escape(text),st);w,h=p.wrap(width,1000);p.drawOn(c,x,y-h);return y-h

def build_deck():
 out=DEST/'AetherGrid_Pitch_Deck_Praveen_Sudireddy.pdf';c=canvas.Canvas(str(out),pagesize=(960,540));c.setAuthor(IDENTITY);c.setTitle('AetherGrid — independent personal project');prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5);prs.core_properties.author=IDENTITY;prs.core_properties.title='AetherGrid — Independent Project by Praveen Sudireddy'
 def textbox(slide,text,x,y,w,h,size=16,color=NAVY,bold=False):
  shape=slide.shapes.add_textbox(Inches(x/72),Inches(y/72),Inches(w/72),Inches(h/72));tf=shape.text_frame;tf.word_wrap=True;tf.margin_left=Pt(0);tf.margin_right=Pt(0);tf.margin_top=Pt(0);tf.margin_bottom=Pt(0)
  for i,line in enumerate(text.split('\n')):
   p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name='DejaVu Sans';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color.lstrip('#'));p.space_after=Pt(9)
  return shape
 for i,(kicker,title,subtitle,lt,left,rt,right) in enumerate(SLIDES,1):
  c.setFillColor(white);c.rect(0,0,960,540,fill=1,stroke=0);c.setFillColor(HexColor(TEAL));c.rect(0,530,960,10,fill=1,stroke=0)
  draw_paragraph(c,'AETHERGRID / '+kicker.upper(),48,505,864,10.5,TEAL,True);draw_paragraph(c,title,48,466,864,27,NAVY,True);draw_paragraph(c,subtitle,48,419,864,11.4,SLATE)
  sl=prs.slides.add_slide(prs.slide_layouts[6]);sl.background.fill.solid();sl.background.fill.fore_color.rgb=RGBColor(255,255,255)
  textbox(sl,'AETHERGRID / '+kicker.upper(),48,35,864,18,10.5,TEAL,True);textbox(sl,title,48,74,864,45,27,NAVY,True);textbox(sl,subtitle,48,121,864,43,11.4,SLATE)
  for x,head,items in [(48,lt,left),(494,rt,right)]:
   c.setFillColor(HexColor(LIGHT));c.roundRect(x,92,418,286,10,fill=1,stroke=0)
   y=draw_paragraph(c,head,x+18,357,382,16,TEAL,True)-17
   for item in items:y=draw_paragraph(c,'• '+item,x+18,y,378,12.2)-13
   if y<100:raise ValueError((i,'PDF overflow',y))
   from pptx.enum.shapes import MSO_SHAPE
   shape=sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x/72),Inches(162/72),Inches(418/72),Inches(286/72));shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string(LIGHT[1:]);shape.line.fill.background()
   textbox(sl,head,x+18,183,382,31,16,TEAL,True);textbox(sl,'\n'.join('• '+v for v in items),x+18,225,378,205,12.2)
  draw_paragraph(c,IDENTITY,48,63,864,8.1,SLATE);draw_paragraph(c,'Independent personal project • Author role only; no employer endorsement • Software-only / simulated outcomes',48,46,820,7.4,SLATE);c.setFont('AG',8);c.drawRightString(912,27,f'{i} / {len(SLIDES)}');c.showPage()
  textbox(sl,IDENTITY,48,477,864,16,8.1,SLATE);textbox(sl,'Independent personal project • Author role only; no employer endorsement • Software-only / simulated outcomes',48,494,820,17,7.4,SLATE);textbox(sl,f'{i} / {len(SLIDES)}',871,511,50,15,8,SLATE)
 c.save();prs.save(DEST/'AetherGrid_Project_Praveen_Sudireddy.pptx')

PAGES=[
 ('AetherGrid: a personal cold-chain intervention project',[
 ('Author and publication status',IDENTITY+'. Independent personal project. The professional role is author context only: this is not presented as an Infosys product, commissioned project, sponsorship or endorsement. This document is a project technical/research report, not a peer-reviewed publication or food-safety approval.'),
 ('Abstract','AetherGrid explores a software loop connecting signed telemetry, a heuristic biological clock, buyer eligibility, time-boxed reservations, manual operator review and deterministic simulated settlement. The current primary implementation uses ASP.NET Core 10, React/TypeScript, SQLite WAL documents and a durable SSE outbox. Python/FastAPI remains a reference and demonstration implementation. Hardware fabrication is excluded. Sensor inputs, funds, insurance calculations and farmer credits in the verified runs are synthetic.'),
 ('Problem and proposed use','An alert alone does not establish who may redirect a load, whether a buyer can receive the complete commodity lot, or how a failed reservation is unwound. The project makes these steps explicit for an illustrative Bengaluru–Hyderabad mango pilot, with 5–10 trucks only as a future supervised scope. No office location, field deployment, buyer partnership or economic benefit is claimed.'),
 ('Classification','Implemented means executable software checked by the supplied tests. Simulated means generated readings and no-money ledger outcomes. Unverified means real-world quality, physical delivery, provider acceptance and operational deployment. Keeping those categories separate is part of the project design.')]),
 ('1 / Architecture, layers and workflow',[
 ('L0–L1: provider boundary and integrity','L0 is a software telemetry provider, not a constructed sensor pod. L1 authenticates users/devices and verifies software HMAC v2, serial/truck/lot bindings, real timestamp freshness and monotonically increasing per-truck sequence. Legacy hsm_serial is a compatibility field name; it does not imply a physical HSM, ECDSA or mTLS.'),
 ('L2: deterministic models','BioTimer accumulates thermal exposure; thermal Dijkstra ranks delivery, buyer liquidation and cold-holding alternatives; eligibility checks full lot, authorization, commodity, capacity, simulated funds, radius, remaining model life and deadline. Settlement is deterministic arithmetic, not insurer adjudication. There is no trained ML model.'),
 ('L3: transaction and persistence','The native service uses SQLite WAL/FULL transactions and JSON documents. BEGIN IMMEDIATE serializes state reads and resource changes. State and sequenced outbox events commit together. Reviewed baseline, checkpoint, cursor, auction, reservation, settlement, evidence snapshot, receipt and sandbox obligation survive restart. Native and Python stores are separate schemas; there is no automatic database migration.'),
 ('L4: five portals and state sequence','React presents Command, Buyer, Analytics, Carrier and Insurance views. Native browser events use SSE with durable cursor replay; the Python demonstration uses Socket.IO. The core sequence is Open → AwaitingConfirmation → Settled, or expiry/rejection → one-time resource release → bounded fallback. A reported full-lot receipt releases capacity once. No endpoint establishes independent physical receipt.')]),
 ('2 / Biological, routing and settlement models',[
 ('Biological clock','Temperature rate is Q10^((T−5)/10), multiplied by implementation-defined humidity and ethylene factors. Exposure increment is rate × interval / reference shelf hours, with a factor of five for door-open intervals. Mango defaults are 264 reference hours, Q10 2.8 and optimal temperature 12°C. These parameters and factors require commodity-specific experimental review. Native state integrates elapsed signed timestamps, persists checkpoints and stops excessive exposure gaps rather than inventing missing history.'),
 ('RVH and routing','Optimistic remaining viable hours at optimal conditions are not the spoilage ETA at the current temperature. Thermal path cost includes edge time, projected temperature, reefer state, cellular-coverage penalty, road/risk factors and life urgency. Reference ranking preserves Python formulas; the operational endpoint adds candidate filters. Optional automatic model routing changes advice only: it does not dispatch a truck, change the original destination or restore consumed life. Fixed graph and straight-line ETA assumptions are not validated navigation.'),
 ('Explicit worked baseline','For a synthetic 1,000 kg mango lot with a ₹1,20,000 contract value and ₹78/kg bid: salvage is ₹78,000; sum assured is ₹72,000; the 45% cover cap is ₹32,400; a 5% deductible gives ₹30,780 calculated insurance. The simulated farmer credit is ₹1,08,780, or 90.65% of contract value. Buyer discount from ₹120/kg is 35%. This corrects earlier unsupported recovery headlines; it is not a guarantee for other policies or lots.'),
 ('Bounded top-up and financial semantics','Insurance is limited by shortfall and configured cover. Any farmer-floor top-up is limited by available modeled penalty; an unfunded universal guarantee is not implemented. Financial obligations use integer paise in the sandbox reconciliation boundary. No bank, UPI rail or insurer is contacted, and no actual money moves.')]),
 ('3 / Security, operations and failure boundaries',[
 ('Identity and authorization','Browser passwords use PBKDF2; opaque HttpOnly sessions carry server-side expiry/version state. Secure cookies require the production flag and TLS. CSRF and same-origin mutation checks protect cookie-based writes. Device and buyer bindings restrict scope; auditors are read-only apart from their own credential maintenance. Acceptance testing found a real buyer-portal authorization mismatch; summary, auctions and settlements were corrected to return that buyer’s scope without exposing other awards.'),
 ('Credential lifecycle','Operator-issued 20-minute recovery tokens are hashed, one-use and manually delivered through a trusted channel. Password change/recovery revokes old sessions. API-key rotation atomically revokes the old digest. Serial-scoped software signing roots are read from an atomically replaced private keyring, with explicit expiry/overlap. No hardware signing is implied.'),
 ('Recovery and retention','SIGKILL with an unconfirmed reservation was tested: a fresh process retained state and released resources once after deadline. Stale telemetry generated a critical readiness alert. Retention removes expired authentication state while preserving samples, evidence, audit, obligations and replay cursors. Online backup must use SQLite’s backup API; old backups may lose newer watermarks and require quarantine/reconciliation before ingestion resumes.'),
 ('Unresolved production boundaries','Single-instance SQLite/document scans are not distributed-scale guarantees. No host-compromise resistance, independent tamper-proof audit archive, external penetration test, SSO/MFA decision, certified TLS deployment or offsite disaster-recovery approval is claimed. A document hash is not proof of sensor truth, consent, insurer acceptance or legal non-repudiation.')]),
 ('4 / Verification protocol and executed results',[
 ('Fresh build and baseline regressions','A fresh npm install/build and locked NuGet restore/Release publish were executed. The Python suite passed 51 tests with no failures/errors/skips. Existing calculation fixtures passed 2,230 assertions. Native ranked routing matched 70 deterministic cases across seven profiles, with 2,100 compared fields. These suites establish tested-case compatibility, not exhaustive whole-server equivalence.'),
 ('Two-role browser acceptance','A fresh isolated native database was provisioned through authenticated APIs. Python-signed telemetry opened a distress auction. An actual buyer browser logged in and reserved the full lot; a separate operator browser confirmed. No settlement existed before confirmation; the UI then showed ₹1,08,780. Another buyer could not read that award, its totals or private batch data. Receipt and duplicate-release behavior were checked through API, alongside CSV reporting and browser logout. No page exceptions or unexpected failed UI requests occurred.'),
 ('Supplementary and adverse checks','Five authenticated portal routes loaded successfully. Python demo pages/APIs, v2 telemetry WebSocket accept/replay rejection and Socket.IO polling were separately smoke-tested. The optional protected calculation API rejected unauthorized calls and matched its settlement fixture. Recovery, rotation, retention, force-kill, rate limiting and bounded concurrency were exercised. Measured local load timings are evidence of this run only, not a production capacity SLA.'),
 ('Dependency and deployment status','Current npm audit reported zero vulnerabilities; the NuGet vulnerable-package query reported none for the native project dependencies. These are snapshot dependency scans, not application security certification. Container recipes/Compose/CI exist, but Docker/Podman runtime access is absent here: no image build/scan or cloud/TLS deployment is asserted.')]),
 ('5 / Experiments, weekly-style journal and no-access lab',[
 ('Historical experiments: distinct evidence','Earlier preserved studies used 1,000 independent eligibility scenarios and 1,000 five-truck software replays. The independent study yielded 440/1,000 eligible reservations. The fleet study contained 5,000 lots, 460 distressed lots and 168 ledger awards: 36.52% of distressed lots. Awarded-only mean recovery was 92.0473%. These synthetic distributions differ from current acceptance tests and are not causal comparisons, field rescue rates or outcomes for all farmers.'),
 ('Progress grouped for weekly reporting','Foundation: FastAPI/React cold-start demo and portal/transport repairs. Correctness: claims audit, explicit payout baseline, eligibility, manual confirmation and bounded fallback. Hardening: v2 access/checkpoints/calculation parity; v3 native workflow/authentication; v4 ranked routing, recovery/rotation/retention and offline lab. Acceptance: current buyer-role defect correction, two-browser journey, artifact identity correction and safe workspace consolidation. These are work-package groupings, not fabricated calendar-week timesheets; exact week dates were not supplied.'),
 ('Local provider emulator','An invented local HTTP financial contract was exercised for stable idempotency, conflicting amounts, ambiguous 502 after commit and an actual read timeout after commit. Tests reconciled by lookup instead of blindly repeating submission. Native callbacks reject stale events, wrong obligations and terminal reversal. This is not a named provider sandbox or a verified live adapter.'),
 ('Synthetic calibration QA','Seven synthetic commodity pipelines require opt-in, reject repeated lots crossing train/holdout, and produce candidate Q10 reports without altering runtime coefficients. Generated error values are pipeline diagnostics, not model accuracy. Real calibration needs controlled independent-lot observations, quality endpoints, provenance, maturity/handling information and expert review.')]),
 ('6 / Reproduction, project identity and remaining gates',[
 ('Reproduction and source map','README.md is the single master Markdown document. Run scripts/build_software.sh with .NET SDK 10 and Node/npm, then configure private users/keys/DB and run the native assembly. scripts/run_v4_checks.py exercises the native regression/adversity paths; scripts/verify_ui_journey.py reproduces the fresh two-role browser flow; scripts/verify_reference_services.py covers supplementary services. Tests and result JSON live under backend/tests and verification/final. The README contains complete commands, API notes, operations and cleanup decisions.'),
 ('Project ownership and use','AetherGrid is Praveen Sudireddy’s independent personal engineering project. Digital Specialist Engineer (DSE), Infosys identifies the author’s professional role only. No employer product, authorization, endorsement, office location, commercial adoption, partnership or investment claim is made. Uses are engineering learning, no-money rehearsal and inspectable decision-support experiments.'),
 ('What remains external','Docker build/image scan must run on a capable host/CI. Deployment security, TLS/proxy behavior, alert delivery, capacity planning and operating ownership need environment-specific sign-off. Live integrations require a selected provider product, permitted test access, actual contract/callback/reconciliation semantics and credentials. Biological/pilot claims require genuine observations, consent, qualified quality review and independently verified outcomes. Hardware fabrication remains excluded.'),
 ('Conclusion','The verified achievement is an inspectable and reproducible software loop with authorization, consistent simulated arithmetic and adverse-path tests. Neither passing tests nor synthetic fit error proves real-world rescue, safe-to-eat quality, actual financial settlement or production readiness.')])]

def build_paper():
 body=ParagraphStyle('body',fontName='AG',fontSize=10.1,leading=14.1,textColor=HexColor(NAVY),spaceAfter=9)
 title=ParagraphStyle('title',parent=body,fontName='AGB',fontSize=19,leading=24,spaceAfter=15)
 head=ParagraphStyle('head',parent=body,fontName='AGB',fontSize=11.5,leading=16,textColor=HexColor(TEAL),spaceAfter=5)
 story=[]
 for i,(name,sections) in enumerate(PAGES):
  if i:story.append(PageBreak())
  story.append(Paragraph(html.escape(name),title))
  for h,text in sections:story.extend([Paragraph(html.escape(h),head),Paragraph(html.escape(text),body)])
 def footer(c,doc):
  c.setFillColor(HexColor(SLATE));c.setFont('AG',7.2);c.drawString(46,38,IDENTITY);c.drawString(46,25,'Independent personal project • No employer endorsement • Software-only prototype');c.drawRightString(550,25,str(doc.page))
 SimpleDocTemplate(str(DEST/'AetherGrid_Paper_Praveen_Sudireddy.pdf'),pagesize=A4,rightMargin=46,leftMargin=46,topMargin=40,bottomMargin=62,title='AetherGrid — Independent Personal Project Technical Report',author=IDENTITY).build(story,onFirstPage=footer,onLaterPages=footer)

if __name__=='__main__':build_deck();build_paper();print('Built personal-project slide PDF, editable PPTX and technical/research PDF.')
