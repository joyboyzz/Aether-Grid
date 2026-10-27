"""Check current personal-project PDFs/PPTX; render editable slides when LibreOffice exists."""
from pathlib import Path
import json,re,shutil,subprocess,os
import pymupdf as fitz
from pptx import Presentation
ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'verification/final/document-qa';OUT.mkdir(parents=True,exist_ok=True)
IDENTITY='Praveen Sudireddy';ROLE='Digital Specialist Engineer (DSE), Infosys'
def norm(s):return re.sub(r'\W','',s).lower()
def check():
 result={}
 for kind,file,count in [('deck','AetherGrid_Pitch_Deck_Praveen_Sudireddy.pdf',17),('paper','AetherGrid_Paper_Praveen_Sudireddy.pdf',7)]:
  doc=fitz.open(ROOT/'paper'/file);issues=[];text='\n'.join(p.get_text() for p in doc)
  for i,page in enumerate(doc):
   for x0,y0,x1,y1,w,*_ in page.get_text('words'):
    if x0<0 or y0<0 or x1>page.rect.width+.5 or y1>page.rect.height+.5:issues.append(dict(page=i+1,word=w))
   if len(page.get_text().strip())<80:issues.append(dict(page=i+1,error='empty'))
  assert IDENTITY in text and norm(ROLE) in norm(text)
  for banned in ['Mangalore','Mangaluru','founder','chief executive','venture capital','startup company']:
   assert banned.lower() not in text.lower(),(file,banned)
  assert len(doc)==count and not issues,(file,issues)
  result[kind]=dict(pages=len(doc),bounds_issues=issues,identity_checked=True,passed=True)
 prs=Presentation(ROOT/'paper/AetherGrid_Project_Praveen_Sudireddy.pptx');assert len(prs.slides)==17
 text='\n'.join(sh.text for sl in prs.slides for sh in sl.shapes if sh.has_text_frame);assert IDENTITY in text and ROLE in text
 result['pptx']=dict(slides=len(prs.slides),editable_text_shapes=sum(sh.has_text_frame for sl in prs.slides for sh in sl.shapes),identity_checked=True)
 office=shutil.which('libreoffice')
 if office:
  dest=OUT/'office';dest.mkdir(exist_ok=True);profile=ROOT/'.acceptance-private/office-profile';profile.parent.mkdir(exist_ok=True)
  p=subprocess.run([office,'-env:UserInstallation='+profile.as_uri(),'--headless','--convert-to','pdf','--outdir',str(dest),str(ROOT/'paper/AetherGrid_Project_Praveen_Sudireddy.pptx')],capture_output=True,text=True,timeout=60);assert p.returncode==0,p.stderr
  doc=fitz.open(dest/'AetherGrid_Project_Praveen_Sudireddy.pdf');issues=[]
  for i,(page,sl) in enumerate(zip(doc,prs.slides)):
   for sh in sl.shapes:
    if sh.has_text_frame and norm(sh.text) not in norm(page.get_text()):issues.append(dict(slide=i+1,missing=sh.text[:120]))
   for x0,y0,x1,y1,w,*_ in page.get_text('words'):
    if x0<0 or y0<0 or x1>page.rect.width+.5 or y1>page.rect.height+.5:issues.append(dict(slide=i+1,bounds=w))
  assert len(doc)==17 and not issues,issues
  result['pptx_render']=dict(engine='LibreOffice',slides=len(doc),text_loss_or_bounds_issues=issues,passed=True)
 else:result['pptx_render']=dict(executed=False,reason='LibreOffice not installed')
 (OUT/'deliverables-qa.json').write_text(json.dumps(result,indent=2));print(json.dumps(result,indent=2));return result
if __name__=='__main__':check()
